import csv
from datetime import date
from io import StringIO
from pathlib import Path

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.beneficiaires.models import Beneficiaire
from apps.comptes.models import Profile
from apps.formations.models import Formation
from apps.referentiels.models import CycleEnquete, Institution

from .models import KoboSoumission
from .services import traiter_soumission

FICHIERS_KOBO = ["suivi_et_satisfaction.xlsx", "satisfaction_institution.xlsx"]


class ExportKoboChoicesTests(TestCase):
    def setUp(self):
        self.institution = Institution.objects.create(libelle="EFTP", type="eftp", region="djibouti")
        CycleEnquete.objects.create(libelle="Cycle 1", date_debut=date(2026, 11, 1), date_fin=date(2026, 11, 30))
        self.beneficiaire = Beneficiaire.objects.create(
            nom="Ali", prenom="Amina", sexe="F", date_naissance=date(1999, 3, 1),
            region="djibouti", institution=self.institution, telephone="77000000",
        )
        Formation.objects.create(
            beneficiaire=self.beneficiaire, domaine="Informatique",
            date_debut=date(2026, 1, 1), date_fin=date(2026, 6, 30),
        )

    def test_ecrit_les_csv_avec_le_bon_format(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command("export_kobo_choices", stdout=StringIO())
                dossier = Path(tmp) / "kobo_forms" / "media"
                for nom in ("beneficiaires.csv", "institutions.csv"):
                    self.assertTrue((dossier / nom).exists(), f"{nom} devrait exister")

                with open(dossier / "beneficiaires.csv", encoding="utf-8") as f:
                    lignes = list(csv.DictReader(f))
                self.assertEqual(lignes[0]["name"], "B-2026-0001")
                self.assertIn("Ali", lignes[0]["label"])
                self.assertEqual(lignes[0]["nom_prenom"], "Amina Ali")
                self.assertEqual(lignes[0]["institution"], "EFTP")
                self.assertEqual(lignes[0]["date_fin_formation"], "2026-06-30")
                self.assertEqual(lignes[0]["telephone"], "77000000")

                with open(dossier / "institutions.csv", encoding="utf-8") as f:
                    lignes = list(csv.DictReader(f))
                self.assertEqual(lignes[0]["name"], str(self.institution.pk))
                self.assertEqual(lignes[0]["label"], "EFTP")


class GenerateKoboXlsformsTests(TestCase):
    def test_genere_les_formulaires_valides(self):
        import tempfile
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command("generate_kobo_xlsforms", stdout=StringIO())
                dossier = Path(tmp) / "kobo_forms"
                for nom in FICHIERS_KOBO:
                    chemin = dossier / nom
                    self.assertTrue(chemin.exists(), f"{nom} devrait exister")
                    classeur = load_workbook(chemin)
                    self.assertEqual(set(classeur.sheetnames), {"survey", "choices", "settings"})
                    noms_champs = [row[1].value for row in classeur["survey"].iter_rows(min_row=2) if row[1].value]
                    self.assertIn("id_beneficiaire" if "institution" not in nom else "institution", noms_champs)

    def test_pyxform_accepte_les_formulaires_generes(self):
        """Le meme moteur de validation que celui utilise par KoboToolbox a l'import."""
        import tempfile

        from pyxform.xls2xform import xls2xform_convert

        with tempfile.TemporaryDirectory() as tmp:
            with override_settings(BASE_DIR=Path(tmp)):
                call_command("generate_kobo_xlsforms", stdout=StringIO())
                dossier = Path(tmp) / "kobo_forms"
                for nom in FICHIERS_KOBO:
                    resultat = xls2xform_convert(str(dossier / nom), str(dossier / (nom + ".xml")), validate=False)
                    self.assertEqual(resultat, [], f"{nom} : avertissements pyxform inattendus")


class TraiterSoumissionTests(TestCase):
    """La logique d'intégration/retraitement partagée par la synchro automatique
    et le bouton « Réessayer » de la page de contrôle."""

    def setUp(self):
        institution = Institution.objects.create(libelle="EFTP", type="eftp", region="djibouti")
        self.beneficiaire = Beneficiaire.objects.create(
            nom="Ali", prenom="Amina", sexe="F", date_naissance=date(1999, 3, 1),
            region="djibouti", institution=institution,
        )

    def test_reessai_reussi_efface_le_message_d_erreur_precedent(self):
        soumission = KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="1", statut="erreur",
            erreur="Bénéficiaire introuvable : B-0000",
            donnees_brutes={
                "selection/id_beneficiaire": self.beneficiaire.pk,
                "selection/date_contact": "2026-01-01",
                "issue_contact": "injoignable",
                "module_suivi/vague": "m3",
            },
        )
        reussite = traiter_soumission(soumission)
        soumission.refresh_from_db()
        self.assertTrue(reussite)
        self.assertEqual(soumission.statut, "integre")
        self.assertEqual(soumission.erreur, "")

    def test_echec_enregistre_le_message_d_erreur(self):
        soumission = KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="2", statut="nouveau",
            donnees_brutes={
                "selection/id_beneficiaire": "B-INCONNU", "module_suivi/vague": "m3",
                "issue_contact": "injoignable",
            },
        )
        reussite = traiter_soumission(soumission)
        soumission.refresh_from_db()
        self.assertFalse(reussite)
        self.assertEqual(soumission.statut, "erreur")
        self.assertIn("B-INCONNU", soumission.erreur)

    def test_soumission_avec_module_satisfaction_cree_aussi_une_satisfaction(self):
        from apps.satisfactions.models import Satisfaction

        CycleEnquete.objects.create(libelle="Cycle 1", date_debut=date(2026, 11, 1), date_fin=date(2026, 11, 30))
        soumission = KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="3", statut="nouveau",
            donnees_brutes={
                "selection/id_beneficiaire": self.beneficiaire.pk,
                "selection/date_contact": "2026-12-01",
                "issue_contact": "joint",
                "module_suivi/vague": "m3",
                "module_suivi/situation": "inactif",
                "bascule/faire_satisfaction": "oui",
                "bascule/cycle": "cycle_1",
                "bascule/note_formation": 4, "bascule/note_formateurs": 5, "bascule/note_contenus": 4,
                "bascule/note_equipements": 4, "bascule/note_conditions": 5,
                "bascule/employabilite": "tout_a_fait", "bascule/recommande": "oui",
            },
        )
        reussite = traiter_soumission(soumission)
        self.assertTrue(reussite)
        satisfaction = Satisfaction.objects.get(beneficiaire=self.beneficiaire)
        self.assertEqual(satisfaction.note_accueil, 5)
        self.assertTrue(satisfaction.recommande)


class SoumissionListViewTests(TestCase):
    def setUp(self):
        self.validateur = User.objects.create_user(username="validateur_test", password="motdepasse123")
        Profile.objects.filter(user=self.validateur).update(role="validateur")
        self.consultation = User.objects.create_user(username="consultation_test", password="motdepasse123")
        Profile.objects.filter(user=self.consultation).update(role="consultation")
        KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="1", statut="erreur",
            erreur="Bénéficiaire introuvable", donnees_brutes={},
        )
        KoboSoumission.objects.create(type_formulaire="suivi", kobo_submission_id="2", statut="integre", donnees_brutes={})

    def test_role_consultation_refuse(self):
        self.client.force_login(self.consultation)
        response = self.client.get(reverse("imports:liste"))
        self.assertEqual(response.status_code, 403)

    def test_role_validateur_voit_la_liste_et_les_compteurs(self):
        self.client.force_login(self.validateur)
        response = self.client.get(reverse("imports:liste"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["nb_total"], 2)
        self.assertEqual(response.context["nb_erreur"], 1)
        self.assertEqual(response.context["nb_integre"], 1)

    def test_filtre_par_statut(self):
        self.client.force_login(self.validateur)
        response = self.client.get(reverse("imports:liste"), {"statut": "erreur"})
        soumissions = list(response.context["soumissions"])
        self.assertEqual(len(soumissions), 1)
        self.assertEqual(soumissions[0].kobo_submission_id, "1")


class SoumissionActionsViewTests(TestCase):
    def setUp(self):
        self.validateur = User.objects.create_user(username="validateur_test2", password="motdepasse123")
        Profile.objects.filter(user=self.validateur).update(role="validateur")
        institution = Institution.objects.create(libelle="EFTP", type="eftp", region="djibouti")
        self.beneficiaire = Beneficiaire.objects.create(
            nom="Ali", prenom="Amina", sexe="F", date_naissance=date(1999, 3, 1),
            region="djibouti", institution=institution,
        )

    def test_reessayer_integre_une_soumission_desormais_valide(self):
        soumission = KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="1", statut="erreur", erreur="Bénéficiaire introuvable",
            donnees_brutes={
                "selection/id_beneficiaire": self.beneficiaire.pk,
                "selection/date_contact": "2026-01-01",
                "issue_contact": "injoignable",
                "module_suivi/vague": "m3",
            },
        )
        self.client.force_login(self.validateur)
        response = self.client.post(reverse("imports:reessayer", args=[soumission.pk]))
        self.assertRedirects(response, reverse("imports:liste"))
        soumission.refresh_from_db()
        self.assertEqual(soumission.statut, "integre")

    def test_marquer_doublon_ecarte_la_soumission(self):
        soumission = KoboSoumission.objects.create(
            type_formulaire="suivi", kobo_submission_id="1", statut="erreur", donnees_brutes={},
        )
        self.client.force_login(self.validateur)
        response = self.client.post(reverse("imports:marquer_doublon", args=[soumission.pk]))
        self.assertRedirects(response, reverse("imports:liste"))
        soumission.refresh_from_db()
        self.assertEqual(soumission.statut, "doublon")
