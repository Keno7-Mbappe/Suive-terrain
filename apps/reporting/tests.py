from datetime import date

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.beneficiaires.models import Beneficiaire
from apps.comptes.models import Profile
from apps.referentiels.models import CycleEnquete, Institution
from apps.satisfactions.models import Satisfaction


class RapportCycleExcelTests(TestCase):
    def setUp(self):
        self.eftp = Institution.objects.create(libelle="EFTP", type="eftp", region="djibouti")
        self.inap = Institution.objects.create(libelle="INAP", type="inap", region="djibouti")
        self.cycle = CycleEnquete.objects.create(
            libelle="Cycle 1", date_debut=date(2026, 11, 1), date_fin=date(2026, 11, 30)
        )
        self.beneficiaire_eftp = Beneficiaire.objects.create(
            nom="Ali", prenom="Amina", sexe="F", date_naissance=date(1999, 3, 1),
            region="djibouti", institution=self.eftp,
        )
        self.beneficiaire_inap = Beneficiaire.objects.create(
            nom="Omar", prenom="Yasin", sexe="M", date_naissance=date(1998, 7, 20),
            region="djibouti", institution=self.inap,
        )
        Satisfaction.objects.create(
            beneficiaire=self.beneficiaire_eftp, cycle=self.cycle,
            note_formation=4, note_formateurs=5, note_contenus=4, note_equipements=3, note_accueil=4,
            amelioration_employabilite="tout_a_fait", recommande=True,
        )
        self.validateur_eftp = User.objects.create_user(username="validateur_eftp", password="motdepasse123")
        Profile.objects.filter(user=self.validateur_eftp).update(role="validateur", institution=self.eftp)
        self.consultation = User.objects.create_user(username="lecture", password="motdepasse123")

    def test_saisie_ne_peut_pas_telecharger(self):
        saisie = User.objects.create_user(username="saisie", password="motdepasse123")
        Profile.objects.filter(user=saisie).update(role="saisie", institution=self.eftp)
        self.client.force_login(saisie)
        response = self.client.get(reverse("reporting:cycle_excel", args=[self.cycle.pk]))
        self.assertEqual(response.status_code, 403)

    def test_validateur_scope_a_son_institution_dans_le_classeur(self):
        self.client.force_login(self.validateur_eftp)
        response = self.client.get(reverse("reporting:cycle_excel", args=[self.cycle.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        import io
        classeur = load_workbook(io.BytesIO(response.content))
        self.assertEqual(
            classeur.sheetnames,
            ["Synthèse", "Bénéficiaires", "Satisfaction", "Satisfaction institutionnelle", "Qualité des données"],
        )

        feuille_beneficiaires = classeur["Bénéficiaires"]
        ids_beneficiaires = [row[0].value for row in feuille_beneficiaires.iter_rows(min_row=2)]
        self.assertIn(self.beneficiaire_eftp.pk, ids_beneficiaires)
        self.assertNotIn(self.beneficiaire_inap.pk, ids_beneficiaires)

    def test_index_accessible_en_consultation(self):
        self.client.force_login(self.consultation)
        response = self.client.get(reverse("reporting:index"))
        self.assertEqual(response.status_code, 200)
