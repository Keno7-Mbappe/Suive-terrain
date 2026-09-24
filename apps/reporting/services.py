from openpyxl import Workbook
from openpyxl.styles import Font

from apps.beneficiaires.models import Beneficiaire
from apps.certifications.models import Certification
from apps.formations.models import Formation
from apps.insertions.models import Insertion
from apps.satisfactions.models import Satisfaction, SatisfactionInstitution
from apps.suivis.models import Suivi


def _entete(feuille, colonnes):
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)


def _taux(numerateur, denominateur):
    return round(numerateur / denominateur * 100, 1) if denominateur else 0


def generer_rapport_cycle(cycle, institution_id=None):
    """Construit un classeur Excel de synthèse pour un cycle d'enquête donné,
    optionnellement restreint à une institution (scoping "saisie"/"validateur")."""
    beneficiaires = Beneficiaire.objects.all()
    if institution_id:
        beneficiaires = beneficiaires.filter(institution_id=institution_id)

    total = beneficiaires.count()
    formes = (
        Formation.objects.filter(beneficiaire__in=beneficiaires, statut_formation="achevee")
        .values("beneficiaire").distinct().count()
    )
    certifies = (
        Certification.objects.filter(beneficiaire__in=beneficiaires)
        .values("beneficiaire").distinct().count()
    )
    inseres = (
        Insertion.objects.filter(beneficiaire__in=beneficiaires)
        .values("beneficiaire").distinct().count()
    )
    nb_suivis = Suivi.objects.filter(beneficiaire__in=beneficiaires).values("beneficiaire").distinct().count()
    satisfactions = Satisfaction.objects.filter(beneficiaire__in=beneficiaires, cycle=cycle)
    satisfactions_institution = SatisfactionInstitution.objects.filter(cycle=cycle)
    if institution_id:
        satisfactions_institution = satisfactions_institution.filter(institution_id=institution_id)
    doublons = Beneficiaire.groupes_doublons(beneficiaires)

    classeur = Workbook()

    synthese = classeur.active
    synthese.title = "Synthèse"
    _entete(synthese, ["Indicateur", "Valeur"])
    for ligne in [
        ("Cycle", cycle.libelle),
        ("Période", f"{cycle.date_debut} au {cycle.date_fin}"),
        ("Bénéficiaires", total),
        ("Formés", formes),
        ("Taux d'achèvement (%)", _taux(formes, total)),
        ("Certifiés", certifies),
        ("Taux de certification (%)", _taux(certifies, total)),
        ("Insérés", inseres),
        ("Taux d'insertion (%)", _taux(inseres, total)),
        ("Bénéficiaires suivis", nb_suivis),
        ("Taux de suivi (%)", _taux(nb_suivis, total)),
        ("Répondants satisfaction (ce cycle)", satisfactions.count()),
        ("Taux de réponse (%)", _taux(satisfactions.count(), total)),
        ("Doublons potentiels détectés", doublons.count()),
    ]:
        synthese.append(ligne)
    synthese.column_dimensions["A"].width = 34

    feuille_beneficiaires = classeur.create_sheet("Bénéficiaires")
    _entete(feuille_beneficiaires, ["ID", "Nom", "Prénom", "Sexe", "Région", "Institution", "Tranche d'âge", "Statut"])
    for b in beneficiaires.select_related("institution"):
        feuille_beneficiaires.append([
            b.id_beneficiaire, b.nom, b.prenom, b.get_sexe_display(), b.get_region_display(),
            str(b.institution), b.tranche_age, b.get_statut_display(),
        ])

    feuille_satisfaction = classeur.create_sheet("Satisfaction")
    _entete(feuille_satisfaction, [
        "Bénéficiaire", "Note formation", "Note formateurs", "Note contenus",
        "Note équipements", "Note accueil", "Note globale", "Recommande",
    ])
    for s in satisfactions.select_related("beneficiaire"):
        feuille_satisfaction.append([
            str(s.beneficiaire), s.note_formation, s.note_formateurs, s.note_contenus,
            s.note_equipements, s.note_accueil, s.note_globale, "Oui" if s.recommande else "Non",
        ])

    feuille_satisfaction_institution = classeur.create_sheet("Satisfaction institutionnelle")
    _entete(feuille_satisfaction_institution, [
        "Institution", "Qualité des données", "Outils de collecte", "Tableaux de bord",
        "Appui technique", "Coordination", "Note globale", "Le dispositif répond-il aux besoins ?",
    ])
    for s in satisfactions_institution.select_related("institution"):
        feuille_satisfaction_institution.append([
            str(s.institution), s.note_qualite_donnees, s.note_outils_collecte, s.note_tableaux_bord,
            s.note_appui_technique, s.note_coordination, s.note_globale, s.get_utilite_dispositif_display(),
        ])

    feuille_qualite = classeur.create_sheet("Qualité des données")
    _entete(feuille_qualite, ["Nom", "Prénom", "Date de naissance", "Occurrences"])
    for d in doublons:
        feuille_qualite.append([d["nom"], d["prenom"], d["date_naissance"], d["occurrences"]])

    return classeur
